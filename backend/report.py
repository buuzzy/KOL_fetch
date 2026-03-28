"""输出报告：生成 BD 可直接用的 Excel/CSV。"""

import csv
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from discovery import KOL
from storage import DiffResult
from config import OUTPUT_DIR


def export_ig_csv(kols: list, filename: str = "") -> str:
    if not filename:
        filename = f"ig_kol_list_{datetime.now().strftime('%Y%m%d')}.csv"
    filepath = os.path.join(OUTPUT_DIR, filename)

    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "序号", "KOL名称", "用户名", "平台", "粉丝数", "帖子数",
            "认证", "分类", "内容方向", "港澳相关度",
            "主页链接", "外部链接", "发现关键词",
        ])
        for i, kol in enumerate(kols, 1):
            writer.writerow([
                i,
                kol.name,
                f"@{kol.username}",
                "Instagram",
                kol.follower_count,
                kol.media_count,
                "是" if kol.is_verified else "否",
                kol.category or "未知",
                ", ".join(kol.content_focus) if kol.content_focus else "待分析",
                kol.hk_relevance_score,
                kol.profile_url,
                kol.external_url or "",
                ", ".join(kol.discovered_via_keywords),
            ])
    return filepath


def export_ig_excel(kols: list, filename: str = "") -> str:
    if not filename:
        filename = f"ig_kol_list_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "IG KOL列表"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="E1306C", end_color="E1306C", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "序号", "KOL名称", "用户名", "平台", "粉丝数", "帖子数",
        "认证", "分类", "内容方向", "港澳相关度",
        "主页链接", "外部链接", "发现关键词",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for i, kol in enumerate(kols, 1):
        row = i + 1
        values = [
            i,
            kol.name,
            f"@{kol.username}",
            "Instagram",
            kol.follower_count,
            kol.media_count,
            "是" if kol.is_verified else "否",
            kol.category or "未知",
            ", ".join(kol.content_focus) if kol.content_focus else "待分析",
            kol.hk_relevance_score,
            kol.profile_url,
            kol.external_url or "",
            ", ".join(kol.discovered_via_keywords),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            if col == 11:
                cell.font = Font(color="0563C1", underline="single")

    col_widths = [6, 25, 20, 12, 12, 10, 8, 15, 20, 10, 40, 30, 20]
    for i, width in enumerate(col_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(64 + (i - 1) % 26 + 1)
        ws.column_dimensions[col_letter].width = width

    ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return filepath


def export_threads_csv(kols: list, filename: str = "") -> str:
    if not filename:
        filename = f"threads_kol_list_{datetime.now().strftime('%Y%m%d')}.csv"
    filepath = os.path.join(OUTPUT_DIR, filename)

    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "序号", "KOL名称", "用户名", "平台", "粉丝数",
            "认证", "内容方向", "港澳相关度",
            "主页链接", "发现关键词",
        ])
        for i, kol in enumerate(kols, 1):
            writer.writerow([
                i,
                kol.name,
                f"@{kol.username}",
                "Threads",
                kol.follower_count,
                "是" if kol.is_verified else "否",
                ", ".join(kol.content_focus) if kol.content_focus else "待分析",
                kol.hk_relevance_score,
                kol.profile_url,
                ", ".join(kol.discovered_via_keywords),
            ])
    return filepath


def export_threads_excel(kols: list, filename: str = "") -> str:
    if not filename:
        filename = f"threads_kol_list_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Threads KOL列表"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "序号", "KOL名称", "用户名", "平台", "粉丝数",
        "认证", "内容方向", "港澳相关度",
        "主页链接", "发现关键词",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for i, kol in enumerate(kols, 1):
        row = i + 1
        values = [
            i,
            kol.name,
            f"@{kol.username}",
            "Threads",
            kol.follower_count,
            "是" if kol.is_verified else "否",
            ", ".join(kol.content_focus) if kol.content_focus else "待分析",
            kol.hk_relevance_score,
            kol.profile_url,
            ", ".join(kol.discovered_via_keywords),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            if col == 9:
                cell.font = Font(color="0563C1", underline="single")

    col_widths = [6, 25, 20, 12, 12, 8, 20, 10, 40, 20]
    for i, width in enumerate(col_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(64 + (i - 1) % 26 + 1)
        ws.column_dimensions[col_letter].width = width

    ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return filepath


def export_csv(kols: list[KOL], filename: str = "") -> str:
    if not filename:
        filename = f"kol_list_{datetime.now().strftime('%Y%m%d')}.csv"
    filepath = os.path.join(OUTPUT_DIR, filename)

    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([
            "序号", "KOL名称", "平台", "订阅数", "视频数",
            "总观看数", "内容方向", "港澳相关度",
            "主页链接", "发现关键词", "国家/地区",
        ])
        for i, kol in enumerate(kols, 1):
            writer.writerow([
                i,
                kol.name,
                "YouTube",
                kol.subscriber_count,
                kol.video_count,
                kol.view_count,
                ", ".join(kol.content_focus) if kol.content_focus else "待分析",
                getattr(kol, "hk_relevance_score", 0),
                kol.profile_url,
                ", ".join(kol.discovered_via_keywords),
                kol.country or "未知",
            ])

    return filepath


def export_excel(kols: list[KOL], filename: str = "") -> str:
    if not filename:
        filename = f"kol_list_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "KOL列表"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [
        "序号", "KOL名称", "平台", "订阅数", "视频数",
        "总观看数", "内容方向", "港澳相关度",
        "主页链接", "发现关键词", "国家/地区",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for i, kol in enumerate(kols, 1):
        row = i + 1
        values = [
            i,
            kol.name,
            "YouTube",
            kol.subscriber_count,
            kol.video_count,
            kol.view_count,
            ", ".join(kol.content_focus) if kol.content_focus else "待分析",
            getattr(kol, "hk_relevance_score", 0),
            kol.profile_url,
            ", ".join(kol.discovered_via_keywords),
            kol.country or "未知",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            if col == 9:
                cell.font = Font(color="0563C1", underline="single")

    col_widths = [6, 25, 10, 12, 10, 15, 20, 10, 45, 20, 12]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.auto_filter.ref = ws.dimensions

    wb.save(filepath)
    return filepath


def export_diff_report(diff: DiffResult, filename: str = "") -> str:
    """导出月度轧差报告。"""
    if not filename:
        filename = f"kol_diff_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = Workbook()

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    ws1 = wb.active
    ws1.title = "新增KOL"
    _write_kol_sheet(ws1, diff.new_kols, header_font, header_fill, green_fill)

    ws2 = wb.create_sheet("粉丝增长TOP")
    _write_growth_sheet(ws2, diff.grown_kols, header_font, header_fill)

    ws3 = wb.create_sheet("不再出现")
    _write_kol_sheet(ws3, diff.lost_kols, header_font, header_fill, red_fill)

    ws4 = wb.create_sheet("摘要")
    ws4.append(["指标", "数值"])
    ws4.append(["上期KOL总数", diff.total_old])
    ws4.append(["本期KOL总数", diff.total_new])
    ws4.append(["新增KOL", len(diff.new_kols)])
    ws4.append(["不再出现", len(diff.lost_kols)])
    ws4.append(["有粉丝增长", len(diff.grown_kols)])

    wb.save(filepath)
    return filepath


def _write_kol_sheet(ws, kols, header_font, header_fill, row_fill):
    headers = ["序号", "KOL名称", "订阅数", "内容方向", "主页链接"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for i, kol in enumerate(kols, 1):
        row = i + 1
        values = [
            i, kol.name, kol.subscriber_count,
            ", ".join(kol.content_focus) if kol.content_focus else "待分析",
            kol.profile_url,
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.fill = row_fill


def _write_growth_sheet(ws, grown_kols, header_font, header_fill):
    headers = ["序号", "KOL名称", "当前订阅", "上期订阅", "增长数", "主页链接"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for i, item in enumerate(grown_kols, 1):
        kol = item["kol"]
        row = i + 1
        values = [
            i, kol.name, kol.subscriber_count,
            item["old_subscribers"], item["growth"],
            kol.profile_url,
        ]
        for col, v in enumerate(values, 1):
            ws.cell(row=row, column=col, value=v)
